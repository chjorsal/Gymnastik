"""Excel-eksport af den beregnede plan (samme format som generation 1.0)."""

from io import BytesIO
from typing import List, Optional

import pandas as pd
from openpyxl.styles import Font, PatternFill

from .values import fmt_time

EXPORT_COLS = [
    "Hold", "Type", "Varighed", "AntalPersoner", "Holdtype", "Alder",
    "OpvisningHal", "OpvisningTid",
    "OpvarmningHal", "OpvarmningStartTid", "OpvarmningSlutTid", "Status",
]
HAL_COLS = ["OpvarmningStartTid", "OpvarmningSlutTid", "Hold", "Holdtype", "Alder", "OpvisningHal", "OpvisningTid"]


def _fill(color: str) -> PatternFill:
    return PatternFill(start_color=color, end_color=color, fill_type="solid")


# Status -> (baggrundsfarve, farv hele rækken?, skrift på status-cellen)
STATUS_STYLES = {
    "KAN IKKE PLANLÆGGES": (_fill("FF4444"), True, Font(color="FFFFFF", bold=True)),
    "MANGLER": (_fill("FFD700"), True, Font(color="000000", bold=True)),
    "OK": (_fill("90EE90"), False, Font(color="000000", bold=True)),
    "IGNORERET": (_fill("D3D3D3"), True, None),
}


_INVALID_SHEET_CHARS = str.maketrans("", "", "[]:*?/\\")


def _sheet_name(raw: str, used: set) -> str:
    """Gyldigt og unikt Excel-arknavn: uden de tegn Excel forbyder, højst 31
    tegn, og med ' (2)', ' (3)' ... hvis navnet (efter afkortning) er brugt."""
    base = str(raw).translate(_INVALID_SHEET_CHARS).strip() or "Ark"
    name, n = base[:31], 2
    while name.lower() in used:
        suffix = f" ({n})"
        name = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(name.lower())
    return name


def _hal_sheet_columns(columns) -> list:
    return [
        c.replace("Opvarmning", "").replace("Opvisning", "Show ")
         .replace("StartTid", "Start").replace("SlutTid", "Slut")
        for c in columns
    ]


def _style_hal_sheet(ws) -> None:
    """Skiftevis farvede rækker, fed header og kolonnebredde efter indhold."""
    light = _fill("FFF9E6")
    for i, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row), start=1):
        if i % 2 == 0:
            for cell in row:
                cell.fill = light
    bold = Font(bold=True)
    for cell in ws[1]:
        cell.font = bold
    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=0)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 3, 40)


def _style_output_sheet(ws) -> None:
    """Farver hver række i Output-arket efter dens status."""
    status_col = next((cell.column for cell in ws[1] if cell.value == "Status"), None)
    if not status_col:
        return
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        status_cell = row[status_col - 1]
        style = STATUS_STYLES.get(status_cell.value)
        if style is None:
            continue
        fill, whole_row, font = style
        for cell in (row if whole_row else [status_cell]):
            cell.fill = fill
        if font is not None:
            status_cell.font = font


def build_workbook(df_all: pd.DataFrame, show_halls: Optional[List[str]] = None) -> BytesIO:
    """show_halls: opvisningshallernes navne i den rækkefølge, deres ark skal
    stå i. Udelades den, bruges rækkefølgen hallerne optræder i data."""
    out = df_all.copy()
    out["OpvisningTid"] = out["OpvisningStart"].apply(fmt_time)
    out["OpvarmningStartTid"] = out["OpvarmningStart"].apply(fmt_time)
    out["OpvarmningSlutTid"] = out["OpvarmningSlut"].apply(fmt_time)
    out["Alder"] = out["ErBarn"].apply(lambda b: "Barn" if b else "Voksen")

    export_cols = [c for c in EXPORT_COLS if c in out.columns]
    hal_cols = [c for c in HAL_COLS if c in out.columns]
    ok_rows = out[out["Status"] == "OK"]

    if show_halls is None:
        show_halls = list(dict.fromkeys(out["OpvisningHal"].dropna()))
    used = {"output", "opvarmningplan", "mangler opvarmning"}

    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        # Ét ark pr. opvisningshal forrest — programmet i rækkefølge, klar
        # til at lægge ind i andre systemer.
        for hal in show_halls:
            hal_rows = out[out["OpvisningHal"] == hal].sort_values("Order")
            sheet_name = _sheet_name(f"Opvisning {hal}", used)
            hal_rows[export_cols].to_excel(writer, index=False, sheet_name=sheet_name)
            _style_hal_sheet(writer.sheets[sheet_name])

        out[export_cols].to_excel(writer, index=False, sheet_name="Output")

        ok_rows[export_cols].sort_values(["OpvarmningHal", "OpvarmningStartTid"]).to_excel(
            writer, index=False, sheet_name="OpvarmningPlan"
        )

        problem = out[out["Status"] == "KAN IKKE PLANLÆGGES"][export_cols]
        if len(problem) > 0:
            problem.to_excel(writer, index=False, sheet_name="Mangler Opvarmning")

        for hal in sorted(ok_rows["OpvarmningHal"].dropna().unique()):
            if not hal:
                continue
            hal_df = ok_rows[ok_rows["OpvarmningHal"] == hal].sort_values("OpvarmningStartTid")[hal_cols]
            hal_df.columns = _hal_sheet_columns(hal_df.columns)
            sheet_name = _sheet_name(hal, used)
            hal_df.to_excel(writer, index=False, sheet_name=sheet_name)
            _style_hal_sheet(writer.sheets[sheet_name])

        _style_output_sheet(writer.sheets["Output"])

    buf.seek(0)
    return buf
